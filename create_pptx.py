from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# カラーパレット
TEAL = RGBColor(45, 106, 122)
TEAL_LIGHT = RGBColor(232, 244, 247)
ACCENT = RGBColor(232, 147, 90)
ACCENT_LIGHT = RGBColor(253, 240, 232)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(44, 62, 80)
GRAY = RGBColor(107, 124, 141)
BG = RGBColor(250, 251, 252)


def add_bg(slide, color=BG):
    """スライド背景色を設定"""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, fill_color, line_color=None):
    """色付き矩形を追加"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 bold=False, color=DARK, alignment=PP_ALIGN.LEFT, font_name="Meiryo"):
    """テキストボックスを追加"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=16, color=DARK):
    """箇条書きを追加"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Meiryo"
        p.space_after = Pt(8)
    return txBox


# =========================================================
# スライド1: タイトルスライド
# =========================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # 白紙
add_bg(slide1, TEAL)

# タイトル
add_text_box(slide1, Inches(1), Inches(1.5), Inches(11), Inches(1.2),
             "今日のクロードコードでの学修の振り返り",
             font_size=40, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)

# サブタイトル
add_text_box(slide1, Inches(1), Inches(3.0), Inches(11), Inches(0.8),
             "2026年3月31日",
             font_size=24, color=WHITE, alignment=PP_ALIGN.CENTER)

# 名前
add_text_box(slide1, Inches(1), Inches(4.0), Inches(11), Inches(0.8),
             "齋藤 史枝（看護系大学教員・10年目）",
             font_size=22, color=WHITE, alignment=PP_ALIGN.CENTER)

# タグ
add_shape(slide1, Inches(3.2), Inches(5.5), Inches(2), Inches(0.5), RGBColor(60, 130, 148))
add_text_box(slide1, Inches(3.2), Inches(5.5), Inches(2), Inches(0.5),
             "Nursing Education", font_size=13, color=WHITE, alignment=PP_ALIGN.CENTER)

add_shape(slide1, Inches(5.6), Inches(5.5), Inches(2.2), Inches(0.5), RGBColor(60, 130, 148))
add_text_box(slide1, Inches(5.6), Inches(5.5), Inches(2.2), Inches(0.5),
             "Simulation Training", font_size=13, color=WHITE, alignment=PP_ALIGN.CENTER)

add_shape(slide1, Inches(8.2), Inches(5.5), Inches(1.8), Inches(0.5), RGBColor(60, 130, 148))
add_text_box(slide1, Inches(8.2), Inches(5.5), Inches(1.8), Inches(0.5),
             "AI Enthusiast", font_size=13, color=WHITE, alignment=PP_ALIGN.CENTER)


# =========================================================
# スライド2: 本日の学修内容（概要）
# =========================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide2)

add_text_box(slide2, Inches(0.8), Inches(0.4), Inches(11), Inches(0.8),
             "本日の学修内容", font_size=32, bold=True, color=TEAL)

# 線
line = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(3), Pt(4))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT
line.line.fill.background()

items = [
    "Claude Codeの基本操作（CLI環境）を体験",
    "GitおよびGitHubの基本操作を学修",
    "強みの整理シートを作成",
    "Claude Code活用プランを策定",
    "自己紹介ページ（HTML）を作成",
    "振り返りレポートのテンプレートを作成",
    "PowerPointファイルのプログラム生成に挑戦",
]

for i, item in enumerate(items):
    y = Inches(1.8) + Inches(i * 0.7)
    # 番号の丸
    circle = slide2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), y, Inches(0.45), Inches(0.45))
    circle.fill.solid()
    circle.fill.fore_color.rgb = TEAL
    circle.line.fill.background()
    add_text_box(slide2, Inches(1.2), y, Inches(0.45), Inches(0.45),
                 str(i + 1), font_size=14, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
    add_text_box(slide2, Inches(1.9), y, Inches(9), Inches(0.45),
                 item, font_size=18, color=DARK)


# =========================================================
# スライド3: 強みの整理
# =========================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide3)

add_text_box(slide3, Inches(0.8), Inches(0.4), Inches(11), Inches(0.8),
             "自分の強みの整理", font_size=32, bold=True, color=TEAL)

line = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(3), Pt(4))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT
line.line.fill.background()

strengths = [
    ("1", "研修の企画・運営力",
     "シミュレーショントレーニングなど実践的な研修を\nゼロから設計し、運営できる「学びの場をつくる力」"),
    ("2", "10年の教育現場経験",
     "看護教育の現場で培った深い経験と知見。\n学生の成長に寄り添い続けてきた観察力と対応力"),
    ("3", "体験型教育のデザイン",
     "座学だけでなく「体験して学ぶ」場をつくれる。\nシミュレーションで現場で使えるスキルを育てる"),
]

for i, (num, title, desc) in enumerate(strengths):
    x = Inches(0.8) + Inches(i * 4)
    # カード背景
    add_shape(slide3, x, Inches(1.8), Inches(3.6), Inches(4.5), TEAL_LIGHT, TEAL)
    # 番号
    circle = slide3.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.2), Inches(2.0), Inches(0.5), Inches(0.5))
    circle.fill.solid()
    circle.fill.fore_color.rgb = TEAL
    circle.line.fill.background()
    add_text_box(slide3, x + Inches(0.2), Inches(2.0), Inches(0.5), Inches(0.5),
                 num, font_size=16, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
    # タイトル
    add_text_box(slide3, x + Inches(0.2), Inches(2.7), Inches(3.2), Inches(0.6),
                 title, font_size=20, bold=True, color=TEAL)
    # 説明
    add_text_box(slide3, x + Inches(0.2), Inches(3.4), Inches(3.2), Inches(2.5),
                 desc, font_size=15, color=GRAY)


# =========================================================
# スライド4: Claude Code活用プラン
# =========================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide4)

add_text_box(slide4, Inches(0.8), Inches(0.4), Inches(11), Inches(0.8),
             "Claude Code 活用プラン", font_size=32, bold=True, color=TEAL)

line = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(3), Pt(4))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT
line.line.fill.background()

plans = [
    ("PLAN 1", "授業・研修の企画をAIで加速",
     "シナリオ案の壁打ち\nルーブリック作成\nアンケート分析"),
    ("PLAN 2", "スケジュール管理をAIで仕組み化",
     "スケジュール案の自動生成\n週次メール下書き\nアジェンダ自動作成"),
    ("PLAN 3", "教材・資料づくりを効率化",
     "ハンドアウト作成\n国試対策問題の生成\n英語論文の要約"),
    ("PLAN 4", "AI動画作成への第一歩",
     "台本作成 → AI音声\n→ AI画像 → 動画編集\n反転授業に活用"),
    ("PLAN 5", "マネジメント力をAIで補強",
     "対応案の相談\n1on1質問リスト生成\n年間計画の策定"),
]

for i, (label, title, desc) in enumerate(plans):
    col = i % 3
    row = i // 3
    x = Inches(0.8) + Inches(col * 4.1)
    y = Inches(1.6) + Inches(row * 2.8)
    # カード
    add_shape(slide4, x, y, Inches(3.7), Inches(2.5), WHITE, TEAL)
    # ラベル
    lbl = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.15), y + Inches(0.15), Inches(1.2), Inches(0.35))
    lbl.fill.solid()
    lbl.fill.fore_color.rgb = ACCENT
    lbl.line.fill.background()
    add_text_box(slide4, x + Inches(0.15), y + Inches(0.15), Inches(1.2), Inches(0.35),
                 label, font_size=11, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
    # タイトル
    add_text_box(slide4, x + Inches(0.15), y + Inches(0.6), Inches(3.4), Inches(0.5),
                 title, font_size=16, bold=True, color=TEAL)
    # 説明
    add_text_box(slide4, x + Inches(0.15), y + Inches(1.1), Inches(3.4), Inches(1.2),
                 desc, font_size=13, color=GRAY)


# =========================================================
# スライド5: 今日の気づき・学び
# =========================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide5)

add_text_box(slide5, Inches(0.8), Inches(0.4), Inches(11), Inches(0.8),
             "今日の気づき・学び", font_size=32, bold=True, color=TEAL)

line = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(3), Pt(4))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT
line.line.fill.background()

learnings = [
    ("CLI環境の理解",
     "Claude Codeはクラウド上のCLI環境で動作する。\nデスクトップ（GUI）は存在せず、コマンドで操作する。\nWordやPowerPointは直接使えないが、\nプログラムでファイルを生成できる。"),
    ("Git/GitHubの活用",
     "git remote, git push, git commitなどの基本操作を学んだ。\nファイルをGitHub経由でWindowsに持ってこれる。\nバージョン管理の重要性を実感した。"),
    ("AIの活用可能性",
     "Claude Codeでプログラミングを活用すれば、\nPowerPointなどのOfficeファイルも生成可能。\nAIは「考える」だけでなく「作る」パートナーになる。"),
]

for i, (title, desc) in enumerate(learnings):
    y = Inches(1.8) + Inches(i * 1.8)
    # カード
    add_shape(slide5, Inches(0.8), y, Inches(11.5), Inches(1.5), ACCENT_LIGHT, ACCENT)
    # タイトル
    add_text_box(slide5, Inches(1.2), y + Inches(0.1), Inches(10.5), Inches(0.5),
                 title, font_size=20, bold=True, color=ACCENT)
    # 説明
    add_text_box(slide5, Inches(1.2), y + Inches(0.55), Inches(10.5), Inches(1.0),
                 desc, font_size=14, color=DARK)


# =========================================================
# スライド6: 実行ロードマップ
# =========================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide6)

add_text_box(slide6, Inches(0.8), Inches(0.4), Inches(11), Inches(0.8),
             "今後の実行ロードマップ", font_size=32, bold=True, color=TEAL)

line = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(3), Pt(4))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT
line.line.fill.background()

roadmap = [
    ("今週", TEAL,
     ["Claudeで研修シナリオを1つ作ってみる",
      "直近のプロジェクトのスケジュール案を生成"]),
    ("今月中", RGBColor(74, 154, 173),
     ["授業資料の下書きをAIで作る習慣をつける",
      "教育動画の台本を1本作ってみる"]),
    ("3ヶ月以内", ACCENT,
     ["AI動画を1本完成させる",
      "マネジメントの相談を定期的にAIでする習慣",
      "成果を振り返り、活用プランをアップデート"]),
]

for i, (period, color, tasks) in enumerate(roadmap):
    x = Inches(0.8) + Inches(i * 4.1)
    # ヘッダ
    header = add_shape(slide6, x, Inches(1.8), Inches(3.7), Inches(0.7), color)
    add_text_box(slide6, x, Inches(1.8), Inches(3.7), Inches(0.7),
                 period, font_size=22, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
    # カード
    add_shape(slide6, x, Inches(2.5), Inches(3.7), Inches(3.5), WHITE, color)
    # タスク
    task_text = "\n\n".join([f"  {t}" for t in tasks])
    add_text_box(slide6, x + Inches(0.2), Inches(2.8), Inches(3.3), Inches(3.0),
                 task_text, font_size=15, color=DARK)


# =========================================================
# スライド7: まとめ
# =========================================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide7, TEAL)

add_text_box(slide7, Inches(1), Inches(1.5), Inches(11), Inches(1),
             "まとめ", font_size=36, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)

summary_items = [
    "Claude Codeを使い、CLI環境でのファイル作成・Git操作を体験した",
    "自分の強みを整理し、AI活用プランとして5つの方向性を策定した",
    "プログラミングによるPowerPoint生成など、AIの「作る力」を実感した",
    "AIは看護教育の現場でも強力なパートナーになりうる",
]

for i, item in enumerate(summary_items):
    y = Inches(3.0) + Inches(i * 0.8)
    add_text_box(slide7, Inches(2), y, Inches(9), Inches(0.6),
                 item, font_size=20, color=WHITE)

add_text_box(slide7, Inches(1), Inches(6.5), Inches(11), Inches(0.5),
             "Created with Claude Code  |  2026.03.31",
             font_size=14, color=RGBColor(180, 220, 230), alignment=PP_ALIGN.CENTER)


# =========================================================
# 保存
# =========================================================
output_path = "/home/user/fumikichi99/今日のワーク/今日のクロードコードでの学修の振り返り.pptx"
prs.save(output_path)
print(f"PowerPointファイルを保存しました: {output_path}")
